# ==========================================================================
# UNAI - Universal Supply-Chain Agent (Cognitive Runtime)
# Copyright (c) 2026 Ravin Angara / Bristlecone. All rights reserved.
# Appends a "How UNAI works — per use case" section (animated flow GIFs) to the
# executive deck. SAFE: pure ADD, idempotent (skip if already present). No
# slide removal / dedup (that caused packaging corruption before).
# ==========================================================================
import os, sys, json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(ROOT, "UNAI_Executive_Briefing.pptx")
flows = json.load(open("/tmp/flows.json"))
NAVY_BG = "0B1020"; TEAL = "22D3AA"; TXT = "E8EDFB"; MUT = "8A97B8"
def C(h): return RGBColor.from_string(h)

p = Presentation(DECK)
BLANK = p.slide_layouts[6] if len(p.slide_layouts) > 6 else p.slide_layouts[0]
SW, SH = p.slide_width, p.slide_height

def already():
    for sl in p.slides:
        for sh in sl.shapes:
            if sh.has_text_frame and "How UNAI works" in sh.text_frame.text:
                return True
    return False
if already():
    print("flow section already present — nothing to do"); sys.exit(0)

def darkslide():
    s = p.slides.add_slide(BLANK)
    bg = s.background; bg.fill.solid(); bg.fill.fore_color.rgb = C(NAVY_BG)
    return s

def txt(s, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, italic=False):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    para = tf.paragraphs[0]; para.alignment = align
    r = para.add_run(); r.text = text; f = r.font
    f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = "Calibri"; f.color.rgb = C(color)
    return tb

def foot(s):
    txt(s, 0.5, 7.06, 12.3, 0.3, "Bristlecone · UNAI Cognitive Runtime · modeled figures — not for external distribution", 9, MUT)

# section divider
s = darkslide()
txt(s, 0.7, 2.4, 12, 1.0, "How UNAI works — per use case", 40, TEAL, bold=True)
txt(s, 0.7, 3.5, 12, 0.6, "One shared brain. Many thin executors. Read once, reason once — then reuse the cognition.", 17, TXT, italic=True)
txt(s, 0.7, 4.2, 12, 0.5, "The animations on the following slides play in Slide Show mode.", 13, MUT)
foot(s)

# one slide per use case
for c in flows:
    s = darkslide()
    n = str(c.get("ratio", "")).split(":")[0] or str(c.get("total", ""))
    txt(s, 0.6, 0.32, 12.2, 0.6, c["name"], 27, TEAL, bold=True)
    txt(s, 0.6, 0.95, 12.2, 0.35, "reads once · reasons once · %d thin executors reuse the shared cognition" % c["total"], 13, MUT)
    gif = os.path.join(ROOT, "UNAI_Flow_%s.gif" % c["key"])
    if os.path.exists(gif):
        s.shapes.add_picture(gif, Inches(0.66), Inches(1.5), width=Inches(12.0))   # height auto (keeps aspect)
    txt(s, 0.6, 6.72, 12.2, 0.4,
        "%s specialist agents → 1 UNAI      −%s%% tokens (modeled)      one read · one plan" % (n, c["savedPct"]),
        15, TXT, bold=True, align=PP_ALIGN.CENTER)
    foot(s)

p.save(DECK)
print("added %d flow slides — deck now %d slides" % (len(flows) + 1, len(p.slides.__iter__.__self__._sldIdLst)))
