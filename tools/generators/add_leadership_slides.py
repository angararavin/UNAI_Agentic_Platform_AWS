# ==========================================================================
# UNAI - Universal Supply-Chain Agent (Cognitive Runtime)
# Copyright (c) 2026 Ravin Angara / Bristlecone. All rights reserved.
#
# PROPRIETARY & CONFIDENTIAL. This file, and the architecture, methods and
# ideas it embodies, are the exclusive property of the copyright holders.
# No part may be copied, reproduced, modified, distributed, reverse-engineered,
# or used to create derivative works without prior written permission.
# Shared under confidentiality; unauthorized use or disclosure is prohibited.
# See LICENSE. Integrity: this file is listed in copyright/MANIFEST.sha256.
# SPDX-License-Identifier: LicenseRef-UNAI-Proprietary   [UNAI-COPYRIGHT v1]
# ==========================================================================
"""Append the 3-slide Bristlecone Leadership Case (differentiation · customer value ·
Bristlecone revenue) to the exec deck, placed just before THE ASK.

SAFE by design: pure ADD + reorder — no slide removal, no drop_rel, no dedup
(those caused packaging corruption before). Idempotent via SKIP-if-present.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(ROOT, "UNAI_Executive_Briefing.pptx")
p = Presentation(DECK); BLANK = p.slide_layouts[0]

TEAL="0E9E80"; TEAL2="22D3AA"; NAVY="14233F"; NAVY2="13203A"; MUT="5D6B8C"
PURPLE="EFE7FF"; AMBER="FBEFD6"; CARD="F3F7FC"; ICE="E2FBF3"; WHITE="FFFFFF"; CORAL="E2574C"; ACC="4f8cff"; PT="7A5B12"
def C(h): return RGBColor.from_string(h)
def txt(s,l,t,w,h,text,size,color,bold=False,font="Calibri",align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,spacing=None):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True
    try: tf.auto_size=None
    except Exception: pass
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0; tf.vertical_anchor=anchor
    for i,ln in enumerate(text.split("\n")):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph(); para.alignment=align
        if spacing: para.line_spacing=spacing
        r=para.add_run(); r.text=ln; f=r.font; f.size=Pt(size); f.bold=bold; f.name=font; f.color.rgb=C(color)
    return tb
def rect(s,l,t,w,h,fill,rounded=True,line=None):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(l),Inches(t),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=C(fill)
    if line: shp.line.color.rgb=C(line); shp.line.width=Pt(1)
    else: shp.line.fill.background()
    shp.shadow.inherit=False; return shp
def kicker(s,label):
    rect(s,0.70,0.62,0.18,0.18,TEAL2,rounded=False); txt(s,0.98,0.50,11,0.35,label,12,TEAL,bold=True)
def header(s,label,title,sub,sub_top=1.78):
    kicker(s,label); txt(s,0.70,0.86,11.93,0.85,title,30,NAVY,bold=True,font="Cambria")
    txt(s,0.70,sub_top,11.93,0.9,sub,14,MUT,spacing=1.05)
def bcard(s,l,t,w,h,head,body):
    rect(s,l,t,w,h,CARD,line="E4EAF3")
    txt(s,l+0.25,t+0.18,w-0.5,0.32,head,13,NAVY,bold=True)
    txt(s,l+0.25,t+0.58,w-0.5,h-0.7,body,11,MUT,spacing=1.06)
def stat(s,l,t,w,big,lab,bg,fg):
    rect(s,l,t,w,1.2,bg); txt(s,l+0.2,t+0.14,w-0.4,0.55,big,26,fg,bold=True,font="Cambria"); txt(s,l+0.2,t+0.72,w-0.4,0.4,lab,10.5,MUT,spacing=1.0)

def slide_diff():
    s=p.slides.add_slide(BLANK)
    header(s,"LEADERSHIP CASE · 1 OF 3","Why we win — a category, not a feature",
      "UNAI is a shared cognitive runtime: cognition built once, thin domain executors on top. It CONSOLIDATES agents "
      "rather than ORCHESTRATING them — the opposite of the MCP / multi-agent pattern everyone else ships.")
    x=0.70
    for big,lab,bg,fg in [("15 → 1","agents (RMA) → one runtime",PURPLE,TEAL),("105 → 21","layer builds (−80%)",ICE,TEAL),
                          ("0","re-mapping (canonical ontology)",CARD,NAVY),("SAP-safe","reads governed lakehouse",AMBER,PT)]:
        stat(s,x,2.85,2.86,big,lab,bg,fg); x+=3.04
    bcard(s,0.70,4.35,5.75,2.35,"Consolidation, not orchestration",
      "N specialist agents collapse to one runtime. Orchestrators / MCP keep N agents plus N² coordination — they "
      "manage the sprawl; UNAI removes it. That is a different category, and a harder thing to copy.")
    bcard(s,6.68,4.35,5.95,2.35,"A structural moat: SAP-compliant + neutral",
      "Reads the governed lakehouse copy in-tenant, read-only first — compliant with SAP API Policy v4/2026 that bars "
      "third-party agents from calling SAP APIs directly. Cloud- and model-neutral, accountable autonomy, one audit "
      "trail. Competitors that hit SAP directly can't match this posture.")

def slide_customer():
    s=p.slides.add_slide(BLANK)
    header(s,"LEADERSHIP CASE · 2 OF 3","What the customer gains — save & become more profitable",
      "The runtime pays back on the biggest AI line (inference is ~85% of enterprise AI budgets) and on build/run cost — "
      "then drives supply-chain outcomes beyond IT.")
    x=0.70
    for big,lab,bg,fg in [("58–80%","fewer tokens / run",PURPLE,TEAL),("~$3.6M/yr","inference saved (Sonnet, real load)",ICE,TEAL),
                          ("~4 mo","payback on migration",CARD,NAVY),("2–5%","margin uplift, 60% faster planning",AMBER,PT)]:
        stat(s,x,2.85,2.86,big,lab,bg,fg); x+=3.04
    bcard(s,0.70,4.35,5.75,2.35,"Inference savings scale with the fleet",
      "Sonnet 4.6 @ 100k agent runs/day ≈ $178k/yr on a light workflow; at real agent-task weight (20–50 LLM calls) "
      "≈ $3.6M/yr, and $10M+ on frontier models at global volume. Prompt caching is additive (up to −90% on the shared prefix).")
    bcard(s,6.68,4.35,5.95,2.35,"Cheaper to build, easy to justify",
      "~80% fewer builds (7 shared layers + a pack per use case, not N×7) → ~$3.6M three-year build/run savings for a "
      "10-capability estate. List price is ~31% of modeled customer value — the customer keeps the majority.")

def slide_revenue():
    s=p.slides.add_slide(BLANK)
    header(s,"LEADERSHIP CASE · 3 OF 3","What Bristlecone makes — the commercial prize",
      "Four reinforcing streams on one platform: subscription + implementation + agent-conversion + managed run. "
      "Land with a use case, expand as packs are added (high net revenue retention).")
    rect(s,0.70,2.85,5.75,3.85,WHITE,line="E4EAF3")
    txt(s,1.05,3.08,5.2,0.4,"Near-term ramp — total revenue (model)",16,NAVY,bold=True,font="Cambria")
    ramp=[("Year 1 · prove · ~10 customers",5),("Year 2 · expand · ~40",22),("Year 3 · scale · ~120",65)]
    y=3.70
    for lab,v in ramp:
        txt(s,1.05,y,3.1,0.3,lab,10.5,MUT)
        rect(s,1.05,y+0.30,4.9*v/65,0.24,TEAL2,rounded=False); txt(s,1.05+4.9*v/65+0.1,y+0.27,1.2,0.3,f"${v}M",11,TEAL,bold=True)
        y+=0.90
    txt(s,1.05,6.25,5.3,0.35,"~$327k average ACV, growing per account as packs are added.",10.5,MUT)
    rect(s,6.68,2.85,5.95,3.85,CARD,line="E4EAF3")
    txt(s,7.0,3.08,5.4,0.4,"The path to $1B (~7 years)",16,NAVY,bold=True,font="Cambria")
    tiles=[("$1B","revenue in ~7 years"),("~800","customers at target"),("~$620M","subscription ARR"),("~$372M","services (SI multiple)")]
    gx=7.0; gy=3.62
    for i,(big,lab) in enumerate(tiles):
        cx=gx+(i%2)*2.70; cy=gy+(i//2)*1.18
        rect(s,cx,cy,2.55,1.03,WHITE)
        txt(s,cx+0.2,cy+0.12,2.2,0.5,big,22,TEAL,bold=True,font="Cambria"); txt(s,cx+0.2,cy+0.62,2.2,0.3,lab,10,MUT,spacing=1.0)
    txt(s,7.0,6.10,5.6,0.5,"The Cognitive Runtime is the wedge: consolidate agent sprawl (conversion), compound through "
        "packs (subscription) and run (managed) — SI strength monetized on a product.",10,MUT,spacing=1.03)

def already_present():
    for s in p.slides:
        t=" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        if "LEADERSHIP CASE" in t: return True
    return False

if already_present():
    print("Leadership Case slides already present — skipping (safe, no duplicates added).")
else:
    slide_diff(); slide_customer(); slide_revenue()
    lst=p.slides._sldIdLst; ids=list(lst)
    diff,cust,rev=ids[-3],ids[-2],ids[-1]
    # place the three, in order, immediately BEFORE the last slide (THE ASK)
    for sid in (diff,cust,rev): lst.remove(sid)
    ask=list(lst)[-1]                     # current last = THE ASK
    ask.addprevious(diff); ask.addprevious(cust); ask.addprevious(rev)
    p.save(DECK)
    print("added 3 Leadership Case slides before THE ASK · total slides:", len(Presentation(DECK).slides))
