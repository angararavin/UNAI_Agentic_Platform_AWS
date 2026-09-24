# ==========================================================================
# UNAI - Universal Supply-Chain Agent (Cognitive Runtime)
# Copyright (c) 2026 Ravin Angara / Bristlecone. All rights reserved.
#
# PROPRIETARY & CONFIDENTIAL. This file, and the architecture, methods and
# ideas it embodies, are the exclusive property of the copyright holders.
# No part may be copied, reproduced, modified, distributed, reverse-engineered,
# or used to create derivative works without prior written permission.
# Shared under confidentiality; unauthorized use or disclosure is prohibited.
# See LICENSE. Integrity: this file is listed in copyright/MANIFEST.sha256.
# SPDX-License-Identifier: LicenseRef-UNAI-Proprietary   [UNAI-COPYRIGHT v1]
# ==========================================================================
"""Rebuild the 3 slides lost during earlier editing (Path to Value, Bristlecone
Revenue, The Ask) on the clean 21-slide base, in the deck's design system, and
assemble the final 24-slide UNAI_Executive_Briefing.pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC="/tmp/clean.pptx"; OUT="UNAI_Executive_Briefing.pptx"
p=Presentation(SRC); BLANK=p.slide_layouts[0]

TEAL="0E9E80"; TEAL2="22D3AA"; NAVY="14233F"; NAVY2="13203A"; MUT="5D6B8C"
PURPLE="EFE7FF"; AMBER="FBEFD6"; CARD="F3F7FC"; ICE="E2FBF3"; WHITE="FFFFFF"; PURPLE_T="7A5B12"
def C(h): return RGBColor.from_string(h)
def txt(s,l,t,w,h,text,size,color,bold=False,font="Calibri",align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,spacing=None):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True
    try: tf.auto_size=None
    except Exception: pass
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0; tf.vertical_anchor=anchor
    for i,ln in enumerate(text.split("\n")):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph(); para.alignment=align
        if spacing: para.line_spacing=spacing
        r=para.add_run(); r.text=ln; f=r.font; f.size=Pt(size); f.bold=bold; f.name=font; f.color.rgb=C(color)
    return tb
def rect(s,l,t,w,h,fill,rounded=True,line=None):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(l),Inches(t),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=C(fill)
    if line: shp.line.color.rgb=C(line); shp.line.width=Pt(1)
    else: shp.line.fill.background()
    shp.shadow.inherit=False; return shp
def kicker(s,label):
    rect(s,0.70,0.62,0.18,0.18,TEAL2,rounded=False); txt(s,0.98,0.50,11,0.35,label,12,TEAL,bold=True)
def header(s,label,title,sub,sub_top=1.78):
    kicker(s,label); txt(s,0.70,0.86,11.93,0.85,title,30,NAVY,bold=True,font="Cambria")
    txt(s,0.70,sub_top,11.93,0.9,sub,14,MUT,spacing=1.05)
def icon(s,l,t,n,d=0.6):
    rect(s,l,t,d,d,NAVY2); txt(s,l,t+0.02,d,d-0.04,n,16,WHITE,bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
def clear_shapes(slide):
    for shp in list(slide.shapes): shp._element.getparent().remove(shp._element)

# ---- redraw the UPGRADE slide (index 4) in place with the FIXED layout ----
s=p.slides[4]; clear_shapes(s)
header(s,"THE UPGRADE","A shared cognitive runtime — cognition once, thin executors on top",
  "The recommended evolution of the architecture: provide perception, memory, ontology, evidence and "
  "planning ONCE per goal; lightweight domain executors sit on top. Like an OS with shared services and "
  "lightweight workers — not N agents each re-doing cognition.", sub_top=2.12)
rect(s,0.70,3.15,5.75,3.55,WHITE,line="E4EAF3")
txt(s,1.05,3.35,5.2,0.4,"Cognitive view — how one agent thinks",16,NAVY,bold=True,font="Cambria")
rows=[("Perception","shared · once",TEAL),("Memory","shared",TEAL),("Reasoning / Planner","plan once",TEAL),
      ("Evidence","per executor",MUT),("Action / Collaboration","per executor",MUT),("Explainability","per executor",MUT)]
y=3.92
for name,tag,col in rows:
    txt(s,1.05,y+0.02,3.3,0.32,name,12.5,NAVY,bold=True,anchor=MSO_ANCHOR.MIDDLE)
    chip=PURPLE if col==MUT else ICE
    rect(s,4.55,y,1.75,0.34,chip)
    txt(s,4.55,y+0.02,1.75,0.30,tag,10.5,(MUT if col==MUT else TEAL),bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    y+=0.445
rect(s,6.68,3.15,5.95,3.55,CARD,line="E4EAF3")
txt(s,7.0,3.35,5.4,0.4,"Platform view — how the system is built",16,NAVY,bold=True,font="Cambria")
flow=["Conversation","Planner","Ontology  (shared meaning)","Knowledge / Evidence",
      "Specialized executors","Optimization / solvers","Enterprise systems · SAP · lakehouse · ServiceNow"]
y=3.84
for i,step in enumerate(flow):
    icon(s,7.0,y,str(i+1),d=0.32)
    bold=step.startswith(("Planner","Ontology"))
    txt(s,7.48,y+0.01,5.0,0.32,step,11,NAVY if bold else MUT,bold=bold,anchor=MSO_ANCHOR.MIDDLE)
    y+=0.355
txt(s,7.0,6.36,5.5,0.28,"Governance is cross-cutting across every layer.",10.5,TEAL,bold=True)
rect(s,0.70,6.88,11.93,0.48,AMBER)
txt(s,0.95,6.90,11.5,0.44,"Honest framing: the seven layers are our reference architecture — not an industry standard "
    "(real ones run 4–10+). This eliminates duplicated retrieval & reasoning; it is not a fixed 1/N token cut.",
    10.5,PURPLE_T,anchor=MSO_ANCHOR.MIDDLE)

# ---------------- SLIDE: PATH TO VALUE ----------------
s=p.slides.add_slide(BLANK)
header(s,"PATH TO VALUE","A staged, low-risk rollout",
  "Prove, then expand, then scale — read-only first, writes gated and human-approved until validated.")
phases=[("1","PHASE 1 · WEEKS 0–6","Prove",
         "Stand up UNAI read-only on one use case (e.g. Supplier Disruption) against your lakehouse; induce the canonical ontology; demo measured value.",PURPLE),
        ("2","PHASE 2 · WEEKS 6–14","Expand",
         "Add 2–3 use cases as tool packs (Spares, RMA); enable gated write-back; wire guardrails, observability and the AI gateway.",ICE),
        ("3","PHASE 3 · QUARTER 2+","Scale",
         "Roll across functions on the shared substrate; self-host the model in-perimeter; new capability becomes a configuration.",CARD)]
x=0.70; w=3.85
for n,ph,big,desc,bg in phases:
    rect(s,x,2.75,w,3.55,bg)
    icon(s,x+0.35,3.05,n,d=0.55)
    txt(s,x+0.35,3.75,w-0.7,0.3,ph,10.5,TEAL,bold=True)
    txt(s,x+0.33,4.05,w-0.7,0.55,big,26,NAVY,bold=True,font="Cambria")
    txt(s,x+0.35,4.75,w-0.7,1.4,desc,12,MUT,spacing=1.06)
    x+=4.03
rect(s,0.70,6.55,11.93,0.5,AMBER)
txt(s,0.95,6.58,11.5,0.44,"Read-only first; writes gated and human-approved until validated. Risk is contained at every step.",11,PURPLE_T,anchor=MSO_ANCHOR.MIDDLE)

# ---------------- SLIDE: BRISTLECONE REVENUE ----------------
s=p.slides.add_slide(BLANK)
header(s,"BRISTLECONE REVENUE","Decomposable path to $1B — subscription + services",
  "Four reinforcing streams on one platform. The Exec · Revenue tab models the build-up; every driver is editable — the figures are an internal model.")
rect(s,0.70,2.90,5.75,3.75,WHITE,line="E4EAF3")
txt(s,1.05,3.12,5.2,0.4,"Four reinforcing revenue streams",16,NAVY,bold=True,font="Cambria")
streams=[("Platform subscription","base + capability packs + connectors + autonomy usage"),
         ("Implementation services","stand-up, ontology induction, integration"),
         ("Agent-conversion services","convert existing agents onto the shared substrate"),
         ("Managed run","operate, monitor and optimize the substrate")]
y=3.70
for name,desc in streams:
    rect(s,1.05,y,0.16,0.16,TEAL2,rounded=False)
    txt(s,1.35,y-0.05,4.9,0.3,name,12.5,NAVY,bold=True)
    txt(s,1.35,y+0.24,4.9,0.4,desc,10.5,MUT,spacing=1.0)
    y+=0.72
rect(s,6.68,2.90,5.95,3.75,CARD,line="E4EAF3")
txt(s,7.0,3.12,5.4,0.4,"A modeled path to the target",16,NAVY,bold=True,font="Cambria")
stats=[("~800","customers on the platform"),("~$620M","subscription ARR"),
       ("~$372M","services (SI multiple)"),("~$1B","total, reached in ~7 years")]
gx=7.0; gy=3.70
for i,(big,lab) in enumerate(stats):
    cx=gx+(i%2)*2.70; cy=gy+(i//2)*1.35
    rect(s,cx,cy,2.55,1.18,WHITE)
    txt(s,cx+0.22,cy+0.14,2.4,0.55,big,26,TEAL,bold=True,font="Cambria")
    txt(s,cx+0.22,cy+0.74,2.4,0.35,lab,10.5,MUT,spacing=1.0)
txt(s,0.70,6.80,11.93,0.4,"High net revenue retention as customers add packs carries the estate toward $1B. Internal model — not a forecast.",10.5,MUT)

# ---------------- SLIDE: THE ASK ----------------
s=p.slides.add_slide(BLANK)
header(s,"THE ASK","Greenlight a 6-week, read-only pilot",
  "One use case, one small team, one clear gate — the substrate is already built.")
asks=[("1","One use case","Supplier Disruption on your Databricks — measured value in six weeks.",PURPLE),
      ("2","One small team","A lead engineer + a supply-chain SME; the substrate is already built.",ICE),
      ("3","One clear gate","Decision to expand based on measured compression, autonomy and ROI.",CARD)]
x=0.70; w=3.85
for n,head_t,desc,bg in asks:
    rect(s,x,2.85,w,3.15,bg)
    icon(s,x+0.35,3.15,n,d=0.55)
    txt(s,x+0.33,3.92,w-0.7,0.45,head_t,19,NAVY,bold=True,font="Cambria")
    txt(s,x+0.35,4.55,w-0.7,1.2,desc,12.5,MUT,spacing=1.08)
    x+=4.03
rect(s,0.70,6.30,11.93,0.75,NAVY)
txt(s,0.95,6.40,11.5,0.55,"UNAI — one AI that does the work of many.\nBristlecone · Research-preview demonstrator available to run live today.",
    12.5,WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE,spacing=1.05)

# ---------------- reorder: path-to-value after messaging(idx18) → pos20; revenue+ask stay last ----------------
lst=p.slides._sldIdLst; ids=list(lst)
path,rev,ask=ids[-3],ids[-2],ids[-1]
lst.remove(path)
list(lst)[18].addnext(path)          # after 'What this means for each leader' (messaging)
# rev & ask already at the end in correct order
p.save(OUT)
print("assembled",OUT,"· slides:",len(Presentation(OUT).slides))
