# ==========================================================================
# UNAI - Universal Supply-Chain Agent (Cognitive Runtime)
# Copyright (c) 2026 Ravin Angara / Bristlecone. All rights reserved.
#
# PROPRIETARY & CONFIDENTIAL. This file, and the architecture, methods and
# ideas it embodies, are the exclusive property of the copyright holders.
# No part may be copied, reproduced, modified, distributed, reverse-engineered,
# or used to create derivative works without prior written permission.
# Shared under confidentiality; unauthorized use or disclosure is prohibited.
# See LICENSE. Integrity: this file is listed in copyright/MANIFEST.sha256.
# SPDX-License-Identifier: LicenseRef-UNAI-Proprietary   [UNAI-COPYRIGHT v1]
# ==========================================================================
"""Executive Presenter — refresh the exec deck IN PLACE from live metrics.

Safe by construction: it only CLEARS + REDRAWS the two runtime slides (upgrade,
savings) that already exist — it never adds or removes slides, so it cannot
corrupt the package (no orphan parts, no duplicate slide names). Numbers on the
savings slide come from coworkers/engine_metrics.json.

    python coworkers/update_deck.py
"""
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DECK = os.path.join(ROOT, "UNAI_Executive_Briefing.pptx")
METRICS = os.path.join(HERE, "engine_metrics.json")

TEAL="0E9E80"; TEAL2="22D3AA"; NAVY="14233F"; NAVY2="13203A"; MUT="5D6B8C"
PURPLE="EFE7FF"; AMBER="FBEFD6"; CARD="F3F7FC"; ICE="E2FBF3"; WHITE="FFFFFF"; CORAL="E2574C"; PT="7A5B12"
def C(h): return RGBColor.from_string(h)
def txt(s,l,t,w,h,text,size,color,bold=False,font="Calibri",align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,spacing=None):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True
    try: tf.auto_size=None
    except Exception: pass
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0; tf.vertical_anchor=anchor
    for i,ln in enumerate(text.split("\n")):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph(); para.alignment=align
        if spacing: para.line_spacing=spacing
        r=para.add_run(); r.text=ln; f=r.font; f.size=Pt(size); f.bold=bold; f.name=font; f.color.rgb=C(color)
    return tb
def rect(s,l,t,w,h,fill,rounded=True,line=None):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(l),Inches(t),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=C(fill)
    if line: shp.line.color.rgb=C(line); shp.line.width=Pt(1)
    else: shp.line.fill.background()
    shp.shadow.inherit=False; return shp
def kicker(s,label):
    rect(s,0.70,0.62,0.18,0.18,TEAL2,rounded=False); txt(s,0.98,0.50,11,0.35,label,12,TEAL,bold=True)
def header(s,label,title,sub,sub_top=1.78):
    kicker(s,label); txt(s,0.70,0.86,11.93,0.85,title,30,NAVY,bold=True,font="Cambria")
    txt(s,0.70,sub_top,11.93,0.9,sub,14,MUT,spacing=1.05)
def icon(s,l,t,n,d):
    rect(s,l,t,d,d,NAVY2); txt(s,l,t+0.02,d,d-0.04,n,16,WHITE,bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
def clear(slide):
    for sh in list(slide.shapes): sh._element.getparent().remove(sh._element)
def stext(slide): return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)

def draw_upgrade(s):
    clear(s)
    header(s,"THE UPGRADE","A shared cognitive runtime — cognition once, thin executors on top",
      "The recommended evolution of the architecture: provide perception, memory, ontology, evidence and "
      "planning ONCE per goal; lightweight domain executors sit on top. Like an OS with shared services and "
      "lightweight workers — not N agents each re-doing cognition.", sub_top=2.12)
    rect(s,0.70,3.15,5.75,3.55,WHITE,line="E4EAF3")
    txt(s,1.05,3.35,5.2,0.4,"Cognitive view — how one agent thinks",16,NAVY,bold=True,font="Cambria")
    rows=[("Perception","shared · once",TEAL),("Memory","shared",TEAL),("Reasoning / Planner","plan once",TEAL),
          ("Evidence","per executor",MUT),("Action / Collaboration","per executor",MUT),("Explainability","per executor",MUT)]
    y=3.92
    for name,tag,col in rows:
        txt(s,1.05,y+0.02,3.3,0.32,name,12.5,NAVY,bold=True,anchor=MSO_ANCHOR.MIDDLE)
        rect(s,4.55,y,1.75,0.34,PURPLE if col==MUT else ICE)
        txt(s,4.55,y+0.02,1.75,0.30,tag,10.5,(MUT if col==MUT else TEAL),bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        y+=0.445
    rect(s,6.68,3.15,5.95,3.55,CARD,line="E4EAF3")
    txt(s,7.0,3.35,5.4,0.4,"Platform view — how the system is built",16,NAVY,bold=True,font="Cambria")
    # NOTE: deliberately NOT called an 'orchestrator' — UNAI is one runtime, not a conductor of many agents.
    flow=["Conversation","Planner","Ontology  (shared meaning)","Knowledge / Evidence",
          "Specialized executors","Optimization / solvers","Enterprise systems · SAP · lakehouse · ServiceNow"]
    y=3.84
    for i,step in enumerate(flow):
        icon(s,7.0,y,str(i+1),0.32)
        bold=step.startswith(("Planner","Ontology"))
        txt(s,7.48,y+0.01,5.0,0.32,step,11,NAVY if bold else MUT,bold=bold,anchor=MSO_ANCHOR.MIDDLE)
        y+=0.355
    txt(s,7.0,6.36,5.5,0.28,"All inside the Cognitive Runtime (UNAI); governance is cross-cutting.",10.5,TEAL,bold=True)
    rect(s,0.70,6.88,11.93,0.48,AMBER)
    txt(s,0.95,6.90,11.5,0.44,"Honest framing: the seven layers are our reference architecture — not an industry standard "
        "(real ones run 4–10+). This eliminates duplicated retrieval & reasoning; it is not a fixed 1/N token cut.",
        10.5,PT,anchor=MSO_ANCHOR.MIDDLE)

def draw_savings(s, m):
    clear(s)
    header(s,"THE SAVINGS","What the upgrade buys — measured on the demonstrator",
      "The cognition base (ontology context + one plan) is paid once per goal instead of once per specialist, and "
      "context is retrieved once then reused from cache. Savings scale with how many specialists share the work.")
    rng = f"{m['minPct']}%–{m['maxPct']}%"
    stats=[(rng,"fewer tokens / run",PURPLE,TEAL),("once","cognition base paid per goal",ICE,TEAL),
           ("1 plan","built per goal · context reused",CARD,NAVY),("+caching","prefix billed once, then cheap reads",AMBER,PT)]
    x=0.70
    for big,lab,bg,fg in stats:
        rect(s,x,2.85,2.86,1.35,bg)
        txt(s,x+0.2,2.98,2.5,0.6,big,30,fg,bold=True,font="Cambria")
        txt(s,x+0.2,3.62,2.5,0.5,lab,11,MUT,spacing=1.0)
        x+=3.04
    txt(s,0.70,4.45,11.9,0.35,"Naive (one agent per specialist)  vs  UNAI shared runtime — tokens / run (modeled by the engine)",12.5,NAVY,bold=True)
    data=[(d["label"],d["naive"],d["runtime"]) for d in m["headline"]]
    mx=max([d[1] for d in data]+[1]); x0=3.05; barw=8.9; y=4.95
    for name,naive,unai in data:
        txt(s,0.70,y-0.02,2.3,0.3,name,10.5,MUT)
        rect(s,x0,y,barw*naive/mx,0.22,CORAL,rounded=False); txt(s,x0+barw*naive/mx+0.08,y-0.04,1.3,0.3,f"{naive:,}",9.5,CORAL,bold=True)
        rect(s,x0,y+0.30,barw*unai/mx,0.22,TEAL2,rounded=False); txt(s,x0+barw*unai/mx+0.08,y+0.26,1.3,0.3,f"{unai:,}",9.5,TEAL,bold=True)
        y+=0.72
    b=m.get("benchmark")
    bench = (f"Benchmarked by token_benchmark.py (disruption flow: {b['gen1']:,} → {b['unai']:,} tokens, −{b['savedPct']}%; "
             "prompt caching cuts billed input further). " if b else
             "Benchmarked by token_benchmark.py (prompt caching cuts billed input further). ")
    txt(s,0.70,y+0.02,11.9,0.6, bench +
        "Single-capability goals save little — as expected. Genuine cost reduction on the shared prefix comes from "
        "prompt / KV-cache reuse where the model supports it.", 10.5,MUT,spacing=1.05)

def main():
    m = json.load(open(METRICS)) if os.path.exists(METRICS) else {"minPct":58,"maxPct":80,
        "headline":[{"label":"Disruption · 4 agents","naive":15590,"runtime":6600},
                    {"label":"Spares (IBP) · 12 agents","naive":43210,"runtime":11020},
                    {"label":"RMA · 15 agents","naive":50560,"runtime":9960}],"benchmark":None}
    p = Presentation(DECK); n_up=n_sv=0
    for s in p.slides:
        t = stext(s)
        if "cognition once, thin executors" in t: draw_upgrade(s); n_up+=1
        elif "What the upgrade buys" in t: draw_savings(s, m); n_sv+=1
    p.save(DECK)
    print(f"deck refreshed in place · upgrade slides {n_up} · savings slides {n_sv} · savings {m['minPct']}%–{m['maxPct']}% · total slides {len(Presentation(DECK).slides)}")

if __name__ == "__main__":
    main()
